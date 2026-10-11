"""Builds a FICTIONAL master resume used for tests and the document proof.

Every name, employer, and number here is invented. The layout deliberately
exercises formatting the editor must preserve: custom margins, Arial at mixed
sizes, centered header, right-aligned dates via a tab stop, style-based
bullets, and bold technology names inside bullets.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

FONT = "Arial"  # metric-compatible with Liberation Sans in the Linux container

# (text, bold) segments per bullet
Bullet = list[tuple[str, bool]]

EXPERIENCE: list[tuple[str, str, str, list[Bullet]]] = [
    (
        "Data Analyst Intern",
        "Fictional Analytics Co., Springfield",
        "Jun 2025 - Aug 2025",
        [
            [
                ("Built ", False),
                ("Python", True),
                (" and ", False),
                ("SQL", True),
                (" scripts that reconciled weekly sales extracts across three regions.", False),
            ],
            [
                ("Created a ", False),
                ("Power BI", True),
                (" dashboard used by the operations team to track open support tickets.", False),
            ],
            [("Documented data definitions for 40 report fields in a shared glossary.", False)],
        ],
    ),
    (
        "IT Support Assistant",
        "Example State University, Springfield",
        "Sep 2023 - May 2025",
        [
            [("Resolved hardware and account access requests for students and staff.", False)],
            [("Wrote step-by-step guides that reduced repeat password reset questions.", False)],
        ],
    ),
]

PROJECTS: list[tuple[str, str, list[Bullet]]] = [
    (
        "Access Review Tracker",
        "Python, pandas, SQLite",
        [
            [
                ("Modeled a sample user-access review as a ", False),
                ("SQLite", True),
                (" database with role and approval tables.", False),
            ],
            [
                ("Flagged accounts without a recorded approval using ", False),
                ("pandas", True),
                (" checks.", False),
            ],
            [("Summarized findings in a one-page control testing memo.", False)],
        ],
    ),
]


def _set_run_font(run, size: float, bold: bool = False) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold or None
    # Set the East Asian slot too, otherwise Word may fall back to another font.
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)


def _paragraph(document, space_after: float = 0, align=None):
    p = document.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = 1.0
    if align is not None:
        p.alignment = align
    return p


def _heading(document, text: str) -> None:
    p = _paragraph(document, space_after=2)
    p.paragraph_format.space_before = Pt(6)
    _set_run_font(p.add_run(text.upper()), 11, bold=True)


def _title_line(document, left_bold: str, left_rest: str, right: str, right_tab: float) -> None:
    p = _paragraph(document)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(right_tab), WD_TAB_ALIGNMENT.RIGHT)
    _set_run_font(p.add_run(left_bold), 10.5, bold=True)
    _set_run_font(p.add_run(left_rest), 10.5)
    _set_run_font(p.add_run("\t" + right), 10.5)


def _bullet(document, segments: Bullet) -> None:
    p = document.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    for text, bold in segments:
        _set_run_font(p.add_run(text), 10, bold=bold)


def build_fixture_master(path: Path) -> Path:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Inches(0.6))
    section.top_margin = section.bottom_margin = Inches(0.5)
    right_tab = 8.5 - 1.2

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)

    name = _paragraph(document, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_run_font(name.add_run("Jordan Example"), 16, bold=True)
    contact = _paragraph(document, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_run_font(
        contact.add_run(
            "Springfield, ST | jordan@example.com | (555) 010-0199 | example.com/jordan"
        ),
        9.5,
    )

    _heading(document, "Education")
    _title_line(
        document, "Example State University", ", B.S. Information Systems", "May 2025", right_tab
    )

    _heading(document, "Experience")
    for title, employer, dates, bullets in EXPERIENCE:
        _title_line(document, title, f" | {employer}", dates, right_tab)
        for segments in bullets:
            _bullet(document, segments)

    _heading(document, "Projects")
    for project, tech, bullets in PROJECTS:
        p = _paragraph(document)
        _set_run_font(p.add_run(project), 10.5, bold=True)
        _set_run_font(p.add_run(" | "), 10.5)
        _set_run_font(p.add_run(tech), 10.5, bold=True)
        for segments in bullets:
            _bullet(document, segments)

    _heading(document, "Skills")
    for label, value in (
        ("Technical: ", "SQL, Python, pandas, Power BI, Excel, SQLite"),
        ("Languages: ", "English (native), Spanish (professional working)"),
    ):
        p = _paragraph(document)
        _set_run_font(p.add_run(label), 10, bold=True)
        _set_run_font(p.add_run(value), 10)

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path
