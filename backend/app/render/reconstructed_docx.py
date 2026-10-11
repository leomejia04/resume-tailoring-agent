"""A clean, deterministic one-page layout for PDF masters or unsupported DOCX layouts.

This does not try to imitate the original design; the UI labels it as a
reconstructed layout that the user reviews before approving.
"""

from __future__ import annotations

import io
from dataclasses import dataclass

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from app.render.format import FormatPreferences, is_entry_line
from app.tailoring.models import DraftContent

FONT = "Arial"  # metric-compatible with Liberation Sans in the converter image
SECTION_ORDER = [
    "summary",
    "education",
    "experience",
    "projects",
    "leadership",
    "certifications",
    "awards",
    "skills",
]
SECTION_TITLES = {
    "summary": "Summary",
    "education": "Education",
    "experience": "Experience",
    "projects": "Projects",
    "leadership": "Leadership & Activities",
    "certifications": "Certifications",
    "awards": "Awards",
    "skills": "Skills",
}


@dataclass(frozen=True)
class ExtraEntry:
    """Always-shown context the draft may not cite (e.g. education)."""

    section: str
    entry: str
    date_text: str | None


def _font(run, size: float, bold: bool = False) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold or None
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)


def _paragraph(document, space_after: float = 0, align=None):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.0
    if align is not None:
        paragraph.alignment = align
    return paragraph


def render_reconstructed(
    header_lines: list[str],
    content: DraftContent,
    prefs: FormatPreferences,
    extras: list[ExtraEntry],
) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(0.6)
    section.top_margin = section.bottom_margin = Inches(0.5)
    right_tab = Inches(8.5 - 1.2)
    document.styles["Normal"].font.name = FONT
    document.styles["Normal"].font.size = Pt(10)

    for i, line in enumerate(header_lines[:3]):
        p = _paragraph(document, space_after=0 if i == 0 else 2, align=WD_ALIGN_PARAGRAPH.CENTER)
        _font(p.add_run(line), 16 if i == 0 else 9.5, bold=i == 0)
    if prefs.show_headline and content.headline:
        p = _paragraph(document, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
        _font(p.add_run(content.headline), 11, bold=True)

    by_section: dict[str, list] = {}
    for entry in content.entries:
        by_section.setdefault(entry.section or "experience", []).append(entry)
    drafted = {(e.section, e.entry) for e in content.entries}
    extra_by_section: dict[str, list[ExtraEntry]] = {}
    for extra in extras:
        if (extra.section, extra.entry) not in drafted:
            extra_by_section.setdefault(extra.section, []).append(extra)

    def heading(title: str) -> None:
        p = _paragraph(document, space_after=2)
        p.paragraph_format.space_before = Pt(6)
        _font(p.add_run(title.upper()), 11, bold=True)

    def entry_line(text: str | None, dates: str | None) -> None:
        p = _paragraph(document)
        p.paragraph_format.tab_stops.add_tab_stop(right_tab, WD_TAB_ALIGNMENT.RIGHT)
        _font(p.add_run(text or ""), 10.5, bold=True)
        if dates:
            _font(p.add_run("\t" + dates), 10.5)

    if content.summary:
        heading(SECTION_TITLES["summary"])
        _font(_paragraph(document).add_run(content.summary), 10)

    for name in SECTION_ORDER:
        if name in ("summary", "skills"):
            continue
        entries = by_section.get(name, [])
        extras_here = extra_by_section.get(name, [])
        if not entries and not extras_here:
            continue
        heading(SECTION_TITLES.get(name, name.title()))
        for extra in extras_here:
            entry_line(extra.entry, extra.date_text)
        for entry in entries:
            entry_line(entry.entry, entry.date_text)
            cap = prefs.project_bullets if name == "projects" else prefs.max_bullets_per_entry
            bullets = [b for b in entry.bullets if not is_entry_line(b.text, entry.entry)]
            for bullet in bullets[:cap]:
                p = document.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.0
                _font(p.add_run(bullet.text), 10)

    if content.skills:
        heading(SECTION_TITLES["skills"])
        for skill in content.skills:
            p = _paragraph(document)
            label, sep, rest = skill.text.partition(":")
            if sep and len(label) <= 30:
                _font(p.add_run(label + ":"), 10, bold=True)
                _font(p.add_run(rest), 10)
            else:
                _font(p.add_run(skill.text), 10)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
