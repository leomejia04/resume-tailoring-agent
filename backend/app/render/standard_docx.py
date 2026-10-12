"""The standard layout every resume uses by default (ADR 0011).

Measured from the owner's reference resume (US Letter, Times New Roman,
1.6 cm side and 1.4 cm top/bottom margins): name 17.5 pt bold centered;
contact line 9.5 pt centered with links; section headings 11.5 pt bold
capitals over a 1 pt full-width rule, 15 pt above; body 9.5 pt on an exact
11.75 pt line pitch;
entries as "Organization<TAB>Location" over "Title<TAB>Dates" (single-line
entries as "Entry<TAB>Dates"); round bullets hanging at 0.125 in with text at
0.25 in; terms the master set in bold stay bold; "Label: text" lines (skills,
relevant coursework) with a bold label. Sections appear only when the draft
has content for them, in a fixed order.

Times New Roman is replaced by the metric-compatible Liberation Serif in the
converter image, so line breaks match a Word rendering.
"""

from __future__ import annotations

import io
import re
import uuid
from collections.abc import Mapping, Sequence
from dataclasses import dataclass

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Inches, Pt

from app.documents.docx_edit import TextSegment
from app.documents.extract import HeaderLink
from app.render.format import FormatPreferences, is_entry_line
from app.render.reconstructed_docx import ExtraEntry
from app.render.template_docx import segments_for
from app.tailoring.models import DraftContent, DraftEntry

FONT = "Times New Roman"
NAME_PT = 17.5
HEADING_PT = 11.5
BODY_PT = 9.5
BODY_LINE_PT = 11.75
SIDE_MARGIN = Cm(1.6)
TOP_BOTTOM_MARGIN = Cm(1.4)
TEXT_WIDTH = Emu(Inches(8.5) - 2 * SIDE_MARGIN)
HEADING_BEFORE_PT = 15
ENTRY_BEFORE_PT = 8
BULLET = "•"
LINK_COLOR = "0563C1"
SECTION_ORDER = [
    "summary",
    "education",
    "experience",
    "projects",
    "leadership",
    "certifications",
    "awards",
    "skills",
]
SECTION_TITLES = {
    "summary": "Summary",
    "education": "Education",
    "experience": "Work Experience",
    "projects": "Projects",
    "leadership": "Leadership & Extra-Curricular Activities",
    "certifications": "Certifications",
    "awards": "Awards",
    "skills": "Skills",
}
LABEL_LINE = re.compile(r"^([A-Z][A-Za-z0-9&/(),.' -]{1,40}):\s+(\S.*)$")
# Plain text between two bold terms that stays bold: "Python and Pandas".
_BOLD_CONNECTOR = re.compile(r"^(?:,\s*(?:and\s+)?|\s+and\s+|,\s+)$")
# Pre-pBdr successors in CT_PPr's schema order (so Word accepts the border).
_PBDR_SUCCESSORS = (
    "w:shd",
    "w:tabs",
    "w:suppressAutoHyphens",
    "w:spacing",
    "w:ind",
    "w:contextualSpacing",
    "w:jc",
    "w:outlineLvl",
    "w:rPr",
)


@dataclass(frozen=True)
class EntryDetail:
    organization: str | None
    location: str | None
    title: str | None


def emphasis_segments(text: str, terms: Sequence[str]) -> list[TextSegment]:
    """Bold whole-word occurrences of terms, joining "A, B, and C" into one bold run."""
    segments = segments_for(text, list(terms))
    for i in range(1, len(segments) - 1):
        middle = segments[i]
        if (
            not middle.bold
            and segments[i - 1].bold
            and segments[i + 1].bold
            and _BOLD_CONNECTOR.match(middle.text)
        ):
            segments[i] = TextSegment(middle.text, bold=True)
    merged: list[TextSegment] = []
    for segment in segments:
        if merged and merged[-1].bold == segment.bold:
            merged[-1] = TextSegment(merged[-1].text + segment.text, bold=segment.bold)
        else:
            merged.append(segment)
    return merged


def _run(paragraph, text: str, size: float = BODY_PT, bold: bool = False):  # noqa: ANN001, ANN202
    run = paragraph.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold or None
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    return run


def _paragraph(document, *, before: float = 0, after: float = 0, align=None, body: bool = True):  # noqa: ANN001, ANN202, E501
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    # Body lines use the reference's exact pitch; name and headings use single spacing.
    paragraph.paragraph_format.line_spacing = Pt(BODY_LINE_PT) if body else 1.0
    if align is not None:
        paragraph.alignment = align
    return paragraph


def _plain_bullets(document) -> None:  # noqa: ANN001
    """Make the template's bullet lists draw a small round bullet in the body font.

    The default template uses Symbol's private-use bullet, which LibreOffice
    draws as a large dot; the reference uses an ordinary bullet character.
    """
    numbering = document.part.numbering_part.element
    for level in numbering.iter(qn("w:lvl")):
        number_format = level.find(qn("w:numFmt"))
        if number_format is None or number_format.get(qn("w:val")) != "bullet":
            continue
        level.find(qn("w:lvlText")).set(qn("w:val"), BULLET)
        props = level.find(qn("w:rPr"))
        fonts = props.find(qn("w:rFonts")) if props is not None else None
        if fonts is not None:
            for key in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                fonts.set(qn(key), FONT)
            fonts.attrib.pop(qn("w:hint"), None)


def _bottom_rule(paragraph) -> None:  # noqa: ANN001
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "8"), ("w:space", "1"), ("w:color", "000000")):
        bottom.set(qn(key), value)
    border.append(bottom)
    paragraph._p.get_or_add_pPr().insert_element_before(border, *_PBDR_SUCCESSORS)


def _hyperlink(paragraph, text: str, url: str) -> None:  # noqa: ANN001
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(key), FONT)
    props.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK_COLOR)
    props.append(color)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(int(BODY_PT * 2)))
    props.append(size)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text_element = OxmlElement("w:t")
    text_element.text = text
    text_element.set(qn("xml:space"), "preserve")
    run.append(text_element)
    link.append(run)
    paragraph._p.append(link)


def _contact_line(paragraph, line: str, links: Sequence[HeaderLink]) -> None:  # noqa: ANN001
    """The line's text, with any linked phrase written as a real hyperlink."""
    cursor = 0
    spans = []
    for link in links:
        index = line.find(link.text, cursor)
        if index != -1:
            spans.append((index, index + len(link.text), link.url))
    spans.sort()
    for start, end, url in spans:
        if start < cursor:
            continue
        if start > cursor:
            _run(paragraph, line[cursor:start])
        _hyperlink(paragraph, line[start:end], url)
        cursor = end
    if cursor < len(line):
        _run(paragraph, line[cursor:])


def _label_line(paragraph, text: str) -> None:  # noqa: ANN001
    match = LABEL_LINE.match(text)
    if match is None:
        _run(paragraph, text)
        return
    _run(paragraph, match.group(1) + ":", bold=True)
    _run(paragraph, " " + match.group(2))


def render_standard(
    *,
    header_lines: list[str],
    header_links: Sequence[HeaderLink],
    content: DraftContent,
    prefs: FormatPreferences,
    extras: Sequence[ExtraEntry],
    details: Mapping[tuple[str | None, str | None], EntryDetail],
    bold_terms: Mapping[uuid.UUID, Sequence[str]],
    project_terms: Sequence[str] = (),
) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = SIDE_MARGIN
    section.top_margin = section.bottom_margin = TOP_BOTTOM_MARGIN
    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    _plain_bullets(document)
    right_tab = TEXT_WIDTH

    if header_lines:
        name = _paragraph(document, align=WD_ALIGN_PARAGRAPH.CENTER, body=False)
        _run(name, header_lines[0], NAME_PT, bold=True)
    for line in header_lines[1:3]:
        _contact_line(_paragraph(document, align=WD_ALIGN_PARAGRAPH.CENTER), line, header_links)
    if prefs.show_headline and content.headline:
        headline = _paragraph(document, before=2, align=WD_ALIGN_PARAGRAPH.CENTER, body=False)
        _run(headline, content.headline, HEADING_PT, bold=True)

    first_heading = True

    def heading(title: str) -> None:
        nonlocal first_heading
        p = _paragraph(
            document, before=4 if first_heading else HEADING_BEFORE_PT, after=3, body=False
        )
        p.paragraph_format.keep_with_next = True
        _run(p, title.upper(), HEADING_PT, bold=True)
        _bottom_rule(p)
        first_heading = False

    def header_line(left: str, right: str | None, *, bold_left: bool, before: float) -> None:
        p = _paragraph(document, before=before)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.tab_stops.add_tab_stop(right_tab, WD_TAB_ALIGNMENT.RIGHT)
        _run(p, left, bold=bold_left)
        if right:
            _run(p, "\t" + right)

    def entry_header(name: str, entry: str | None, date_text: str | None, first: bool) -> None:
        before = 0 if first else ENTRY_BEFORE_PT
        detail = details.get((name, entry))
        if detail and detail.organization:
            header_line(detail.organization, detail.location, bold_left=True, before=before)
            header_line(detail.title or "", date_text, bold_left=name == "education", before=0)
        else:
            header_line(entry or "", date_text, bold_left=True, before=before)

    def bullet(entry: DraftEntry, text: str, evidence_ids: Sequence[uuid.UUID]) -> None:
        terms = [t for e in evidence_ids for t in bold_terms.get(e, ())]
        if entry.section == "projects" and prefs.bold_project_technologies:
            terms += list(project_terms)
        p = document.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.125)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(BODY_LINE_PT)
        for segment in emphasis_segments(text, terms):
            _run(p, segment.text, bold=segment.bold)

    by_section: dict[str, list[DraftEntry]] = {}
    for entry in content.entries:
        by_section.setdefault(entry.section or "experience", []).append(entry)
    drafted = {(e.section, e.entry) for e in content.entries}
    extra_by_section: dict[str, list[ExtraEntry]] = {}
    for extra in extras:
        if (extra.section, extra.entry) not in drafted:
            extra_by_section.setdefault(extra.section, []).append(extra)

    if content.summary:
        heading(SECTION_TITLES["summary"])
        _run(_paragraph(document), content.summary)

    for name in SECTION_ORDER:
        if name in ("summary", "skills"):
            continue
        entries = by_section.get(name, [])
        extras_here = extra_by_section.get(name, [])
        if not entries and not extras_here:
            continue
        heading(SECTION_TITLES[name])
        first = True
        for extra in extras_here:
            entry_header(name, extra.entry, extra.date_text, first)
            first = False
        for entry in entries:
            entry_header(name, entry.entry, entry.date_text, first)
            first = False
            for item in entry.bullets:
                if is_entry_line(item.text, entry.entry):
                    continue
                if name == "education" and LABEL_LINE.match(item.text):
                    _label_line(_paragraph(document), item.text)
                else:
                    bullet(entry, item.text, item.evidence_ids)

    if content.skills:
        heading(SECTION_TITLES["skills"])
        for skill in content.skills:
            _label_line(_paragraph(document), skill.text)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
