from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from app.documents.docx_edit import (
    StaleEditError,
    TextSegment,
    UnsupportedLayoutError,
    body_paragraphs,
    edit_copy,
    list_paragraphs,
    sha256_file,
)
from app.documents.docx_verify import compare_formatting
from app.documents.fixtures import build_fixture_master

NEW_BULLET = [
    TextSegment("Automated weekly sales reconciliation with "),
    TextSegment("Python", bold=True),
    TextSegment(" and "),
    TextSegment("SQL", bold=True),
    TextSegment(" across three regions."),
]


@pytest.fixture
def master(tmp_path: Path) -> Path:
    return build_fixture_master(tmp_path / "master.docx")


def first_bullet(path: Path):
    refs = list_paragraphs(Document(str(path)))
    return next(r for r in refs if r.is_bullet)


def test_fixture_bullets_are_detected(master):
    refs = list_paragraphs(Document(str(master)))
    bullets = [r for r in refs if r.is_bullet]
    assert len(bullets) == 8
    assert not refs[0].is_bullet  # name line


def test_edit_preserves_master_and_formatting(master, tmp_path):
    target = first_bullet(master)
    master_hash = sha256_file(master)
    edited = tmp_path / "edited.docx"

    edit_copy(master, edited, target.index, target.text, NEW_BULLET)

    assert sha256_file(master) == master_hash
    report = compare_formatting(master, edited, target.index)
    assert report.ok, report
    assert report.target_run_formats == [
        "Arial 10pt",
        "Arial 10pt bold",
        "Arial 10pt",
        "Arial 10pt bold",
        "Arial 10pt",
    ]
    new_text = body_paragraphs(Document(str(edited)))[target.index].text
    assert new_text == "".join(s.text for s in NEW_BULLET)


def test_bold_derived_when_paragraph_has_no_bold_run(master, tmp_path):
    refs = list_paragraphs(Document(str(master)))
    plain = next(r for r in refs if r.is_bullet and r.text.startswith("Documented"))
    edited = tmp_path / "edited.docx"
    segments = [TextSegment("Documented "), TextSegment("40", bold=True), TextSegment(" fields.")]

    edit_copy(master, edited, plain.index, plain.text, segments)

    report = compare_formatting(master, edited, plain.index)
    assert report.ok
    assert report.target_run_formats == ["Arial 10pt", "Arial 10pt bold", "Arial 10pt"]


def test_stale_expected_text_is_rejected(master, tmp_path):
    target = first_bullet(master)
    with pytest.raises(StaleEditError):
        edit_copy(master, tmp_path / "e.docx", target.index, "old text", NEW_BULLET)


def test_output_cannot_overwrite_master(master):
    target = first_bullet(master)
    with pytest.raises(ValueError):
        edit_copy(master, master, target.index, target.text, NEW_BULLET)


def test_empty_replacement_is_rejected(master, tmp_path):
    target = first_bullet(master)
    with pytest.raises(ValueError):
        edit_copy(master, tmp_path / "e.docx", target.index, target.text, [TextSegment("  ")])


def test_hyperlink_paragraph_is_unsupported(master, tmp_path):
    document = Document(str(master))
    paragraph = body_paragraphs(document)[1]  # contact line
    paragraph._p.append(paragraph._p.makeelement(qn("w:hyperlink"), {}))
    linked = tmp_path / "linked.docx"
    document.save(str(linked))

    with pytest.raises(UnsupportedLayoutError, match="hyperlink"):
        edit_copy(linked, tmp_path / "e.docx", 1, paragraph.text, [TextSegment("x")])


def test_verifier_detects_formatting_change_elsewhere(master, tmp_path):
    target = first_bullet(master)
    edited = tmp_path / "edited.docx"
    edit_copy(master, edited, target.index, target.text, NEW_BULLET)

    document = Document(str(edited))
    body_paragraphs(document)[0].runs[0].font.size = Pt(18)  # tamper with the name line
    document.sections[0].left_margin = Pt(20)
    document.save(str(edited))

    report = compare_formatting(master, edited, target.index)
    assert not report.ok
    assert report.section_changes
    assert any("paragraph 0" in change for change in report.paragraph_changes)
