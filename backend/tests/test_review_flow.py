"""Milestone 5 end-to-end tests: render, preview, edit, revise, approve, download.

Needs the local Supabase stack and the Docker converter image; real PDFs are
rendered and inspected. No real model is called.
"""

from __future__ import annotations

import hashlib

import httpx
import pytest
from docx import Document
from docx.shared import Pt

from app.ai.client import LLMResult
from app.ai.pricing import Usage
from app.ai.schemas import BulletOut, ChatPatchOut, EntryOut
from app.config import Settings
from app.documents.convert import resolve_backend
from app.documents.fixtures import build_fixture_master
from tests.test_isolation import _settings_or_skip, new_account, reset_rate_limits
from tests.test_tailoring_flow import create_job, make_app, prepared_account

pytestmark = [
    pytest.mark.integration,
    pytest.mark.skipif(resolve_backend() is None, reason="PDF converter not available"),
]


@pytest.fixture(scope="module")
def base_settings() -> Settings:
    settings = _settings_or_skip()
    reset_rate_limits(settings)
    return settings


@pytest.fixture(scope="module")
def demo_app(base_settings):
    return make_app(base_settings, llm_provider="demo")


@pytest.fixture(scope="module")
def docx_bytes(tmp_path_factory) -> bytes:
    return build_fixture_master(tmp_path_factory.mktemp("m5") / "master.docx").read_bytes()


def drafted(app, docx_bytes, prefix="JordanExample"):
    account = prepared_account(app, docx_bytes)
    profile = account.client.get("/api/v1/profile").json()
    account.client.put(
        "/api/v1/profile",
        json={
            "expected_version": profile["lock_version"],
            "display_name": "Jordan Example",
            "filename_prefix": prefix,
        },
    )
    job_id = create_job(account)
    run = account.client.post("/api/v1/runs", json={"job_id": job_id}).json()
    run = account.client.post(f"/api/v1/runs/{run['id']}/draft").json()
    assert run["status"] == "awaiting_review", run
    account.run_url = f"/api/v1/runs/{run['id']}"  # type: ignore[attr-defined]
    return account, run


def render(account):
    response = account.client.post(f"{account.run_url}/render")
    assert response.status_code == 201, response.text
    return response.json()


def test_render_preview_approve_download(demo_app, docx_bytes):
    account, run = drafted(demo_app, docx_bytes)
    rendered = render(account)
    artifact = rendered["artifact"]
    assert artifact["passed"] is True, artifact["checks"]
    assert artifact["layout"] == "template" and artifact["page_count"] == 1
    assert artifact["filename"] == "JordanExample_Data_Analyst_Fictional_Retail_Co_CV.pdf"
    assert artifact["draft_version"] == rendered["draft"]["version"]

    pdf = account.client.get(f"/api/v1/artifacts/{artifact['id']}/pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert hashlib.sha256(pdf.content).hexdigest() == artifact["sha256"]
    assert "frame-ancestors 'self'" in pdf.headers["content-security-policy"]
    assert pdf.headers["x-frame-options"] == "SAMEORIGIN"
    assert "inline" in pdf.headers["content-disposition"]
    png = account.client.get(f"/api/v1/artifacts/{artifact['id']}/preview.png")
    assert png.status_code == 200 and png.content.startswith(b"\x89PNG")

    approve = {
        "draft_version": artifact["draft_version"],
        "artifact_id": artifact["id"],
        "artifact_sha256": "0" * 64,
    }
    assert (
        account.client.post(f"{account.run_url}/approve", json=approve).status_code == 409
    )  # wrong hash
    approve["artifact_sha256"] = artifact["sha256"]
    approved = account.client.post(f"{account.run_url}/approve", json=approve)
    assert approved.status_code == 200 and approved.json()["status"] == "approved"
    assert approved.json()["approval"]["artifact_id"] == artifact["id"]
    assert (
        account.client.post(f"{account.run_url}/approve", json=approve).status_code == 409
    )  # already

    approvals = account.client.get("/api/v1/approvals").json()
    assert len(approvals) == 1 and approvals[0]["filename"] == artifact["filename"]
    download = account.client.get(f"/api/v1/approvals/{approvals[0]['id']}/download")
    assert download.status_code == 200
    assert (
        download.headers["content-disposition"] == f'attachment; filename="{artifact["filename"]}"'
    )
    assert hashlib.sha256(download.content).hexdigest() == artifact["sha256"]


def test_edits_create_versions_and_invalidate_approval(demo_app, docx_bytes):
    account, run = drafted(demo_app, docx_bytes)
    artifact = render(account)["artifact"]
    account.client.post(
        f"{account.run_url}/approve",
        json={
            "draft_version": artifact["draft_version"],
            "artifact_id": artifact["id"],
            "artifact_sha256": artifact["sha256"],
        },
    )
    version = artifact["draft_version"]
    entry = next(
        e for e in run["draft"]["content"]["entries"] if e["experience_type"] == "internship"
    )

    bad = account.client.post(
        f"{account.run_url}/draft/edits",
        json={
            "expected_version": version,
            "operations": [
                {
                    "op": "update_bullet",
                    "entry_ref": entry["ref"],
                    "index": 0,
                    "text": "Built Tableau dashboards that boosted sales 40%.",
                }
            ],
        },
    )
    assert bad.status_code == 422 and bad.json()["error"]["code"] == "unsupported_edit"
    codes = {f["type"] for f in bad.json()["error"]["fields"]}
    assert {"unsupported_term", "unsupported_number"} <= codes or "unsupported_name" in codes

    reworded = entry["bullets"][0]["text"].replace("Built", "Wrote", 1)
    good = account.client.post(
        f"{account.run_url}/draft/edits",
        json={
            "expected_version": version,
            "operations": [
                {"op": "update_bullet", "entry_ref": entry["ref"], "index": 0, "text": reworded}
            ],
        },
    )
    assert good.status_code == 200, good.text
    body = good.json()
    assert body["draft"]["version"] == version + 1
    assert (
        body["status"] == "awaiting_review"
        and body["approval"] is None
        and body["artifact"] is None
    )
    assert body["draft"]["created_by"] == "edit"

    stale = account.client.post(
        f"{account.run_url}/draft/edits",
        json={
            "expected_version": version,
            "operations": [{"op": "remove_bullet", "entry_ref": entry["ref"], "index": 0}],
        },
    )
    assert stale.status_code == 409
    old_approve = account.client.post(
        f"{account.run_url}/approve",
        json={
            "draft_version": version,
            "artifact_id": artifact["id"],
            "artifact_sha256": artifact["sha256"],
        },
    )
    assert (
        old_approve.status_code == 409 and old_approve.json()["error"]["code"] == "stale_approval"
    )
    # The earlier approved PDF stays downloadable.
    approval_id = account.client.get("/api/v1/approvals").json()[0]["id"]
    assert account.client.get(f"/api/v1/approvals/{approval_id}/download").status_code == 200

    restored = account.client.post(
        f"{account.run_url}/draft/restore",
        json={"version": version, "expected_version": version + 1},
    )
    assert restored.status_code == 200
    assert restored.json()["draft"]["version"] == version + 2
    assert restored.json()["draft"]["content"]["entries"] == run["draft"]["content"]["entries"]
    history = account.client.get(f"{account.run_url}/draft/versions").json()
    assert [v["version"] for v in history] == [version + 2, version + 1, version]


def test_selection_ops_and_formatting_preferences(demo_app, docx_bytes):
    account, run = drafted(demo_app, docx_bytes)
    version = run["draft"]["version"]
    entry = next(e for e in run["draft"]["content"]["entries"] if len(e["bullets"]) >= 2)
    omitted = run["draft"]["content"]["omitted"]
    ops = [
        {"op": "move_bullet", "entry_ref": entry["ref"], "index": 0, "to": 1},
        {"op": "remove_bullet", "entry_ref": entry["ref"], "index": 1},
        {
            "op": "update_format",
            "format": {
                "max_bullets_per_entry": 2,
                "project_bullets": 3,
                "bold_project_technologies": True,
                "languages_on_own_line": True,
                "show_headline": False,
                "locked_entries": [],
            },
        },
    ]
    if omitted:
        ops.append({"op": "add_fact", "evidence_id": omitted[0]["evidence_id"]})
    response = account.client.post(
        f"{account.run_url}/draft/edits", json={"expected_version": version, "operations": ops}
    )
    assert response.status_code == 200, response.text
    content = response.json()["draft"]["content"]
    edited = next(e for e in content["entries"] if e["ref"] == entry["ref"])
    assert edited["bullets"][0]["text"] == entry["bullets"][1]["text"]
    assert len(edited["bullets"]) == len(entry["bullets"]) - 1
    assert content["format"]["max_bullets_per_entry"] == 2

    prefs = account.client.get("/api/v1/profile/preferences").json()
    remembered = account.client.put(
        "/api/v1/profile/preferences",
        json={
            "expected_version": prefs["lock_version"],
            "preferences": {**prefs["preferences"], "project_bullets": 2},
        },
    )
    assert (
        remembered.status_code == 200 and remembered.json()["preferences"]["project_bullets"] == 2
    )
    job_id = create_job(account)
    new_run = account.client.post("/api/v1/runs", json={"job_id": job_id}).json()
    new_run = account.client.post(f"/api/v1/runs/{new_run['id']}/draft").json()
    assert new_run["draft"]["content"]["format"]["project_bullets"] == 2  # remembered with consent


def test_regenerate_and_chat_revisions(monkeypatch, base_settings, docx_bytes):
    app = make_app(base_settings, llm_provider="demo")
    account, run = drafted(app, docx_bytes)
    version = run["draft"]["version"]
    entry = run["draft"]["content"]["entries"][0]
    regenerated = account.client.post(
        f"{account.run_url}/draft/regenerate",
        json={"entry_ref": entry["ref"], "expected_version": version},
    )
    assert regenerated.status_code == 200, regenerated.text
    body = regenerated.json()
    assert body["draft"]["version"] == version + 1 and body["draft"]["created_by"] == "generation"
    assert [e["ref"] for e in body["draft"]["content"]["entries"]] == [
        e["ref"] for e in run["draft"]["content"]["entries"]
    ]

    demo_chat = account.client.post(
        f"{account.run_url}/draft/chat",
        json={"message": "Make it shorter", "expected_version": version + 1},
    )
    assert demo_chat.status_code == 200 and demo_chat.json()["new_version"] is None
    assert "Demo mode" in demo_chat.json()["reply"]

    class ChatClient:
        provider = "openai"
        model = "scripted-chat"

        def parse(self, *, instructions, payload, schema, max_output_tokens, user_hash):
            assert schema is ChatPatchOut and payload["instruction"] == "Emphasize SQL"
            target = payload["entries"][0]
            fact = target["facts"][0]
            patch = ChatPatchOut(
                reply="I emphasized SQL and added a cost-savings bullet.",
                entries=[
                    EntryOut(
                        entry_ref=target["ref"],
                        bullets=[
                            BulletOut(
                                text=fact["text"], evidence_ids=[fact["id"]], requirement_ids=[]
                            ),
                            BulletOut(
                                text="Cut reporting costs by 60% using Snowflake.",
                                evidence_ids=[fact["id"]],
                                requirement_ids=[],
                            ),
                        ],
                    )
                ],
                headline=None,
                change_summary=False,
                summary=None,
            )
            return LLMResult(
                output=patch,
                usage=Usage(input_tokens=1000, output_tokens=200),
                latency_ms=3,
                model=self.model,
            )

    monkeypatch.setattr("app.api.v1.tailoring_routes.client_for", lambda settings: ChatClient())
    chat = account.client.post(
        f"{account.run_url}/draft/chat",
        json={"message": "Emphasize SQL", "expected_version": version + 1},
    )
    assert chat.status_code == 200, chat.text
    result = chat.json()
    assert result["new_version"] == version + 2
    assert "not supported" in result["reply"]
    text = str(result["run"]["draft"]["content"])
    assert "60%" not in text and "Snowflake" not in text
    history = account.client.get(f"{account.run_url}/chat").json()
    assert [m["role"] for m in history] == ["user", "assistant", "user", "assistant"]
    assert history[-1]["draft_version"] == version + 2


def test_long_resume_is_fit_to_one_page_by_dropping_bullets(demo_app, tmp_path):
    path = build_fixture_master(tmp_path / "long.docx")
    document = Document(str(path))
    originals = [p for p in document.paragraphs if p.style.name == "List Bullet"]
    for n, bullet in enumerate(originals):
        for i in range(3):
            new = document.add_paragraph(
                f"Delivered recurring analysis package {n}-{i} covering inventory, staffing, "
                "service levels, and regional sales trends, with written findings reviewed "
                "every week by the operations leads, shared with store managers, archived in "
                "the team drive, and summarized again at each monthly planning meeting.",
                style="List Bullet",
            )
            bullet._p.addnext(new._p)
            new.runs[0].font.size = Pt(10)
    document.save(str(path))
    account, run = drafted(demo_app, path.read_bytes())
    version = run["draft"]["version"]
    fmt = {
        "max_bullets_per_entry": 8,
        "project_bullets": 3,
        "bold_project_technologies": True,
        "languages_on_own_line": True,
        "show_headline": False,
        "locked_entries": [],
    }
    edited = account.client.post(
        f"{account.run_url}/draft/edits",
        json={"expected_version": version, "operations": [{"op": "update_format", "format": fmt}]},
    )
    assert edited.status_code == 200, edited.text
    artifact = render(account)["artifact"]
    assert artifact["page_count"] == 1
    summary = {
        k: artifact["checks"][k]
        for k in (
            "page_count",
            "dropped_to_fit",
            "capped_by_preference",
            "rendered_bullets",
            "notes",
        )
    }
    assert artifact["checks"]["dropped_to_fit"], f"expected bullets to be dropped: {summary}"
    assert artifact["passed"] is True, artifact["checks"]


def test_pdf_master_uses_labeled_reconstructed_layout(demo_app, docx_bytes, tmp_path):
    from app.documents.convert import docx_to_pdf

    docx_path = tmp_path / "m.docx"
    docx_path.write_bytes(docx_bytes)
    pdf_bytes = docx_to_pdf(docx_path, tmp_path / "m.pdf").read_bytes()
    account = new_account(demo_app)
    master = account.client.post(
        "/api/v1/masters", files={"file": ("resume.pdf", pdf_bytes, "application/pdf")}
    ).json()["master"]
    for item in account.client.get("/api/v1/evidence").json():
        kind = item["suggested_type"] if item["suggested_type"] != "unclassified" else "other"
        account.client.post(
            f"/api/v1/evidence/{item['id']}/versions",
            json={
                "expected_version": 1,
                "experience_type": kind,
                "verification_status": "verified",
            },
        )
    job_id = create_job(account)
    run = account.client.post(
        "/api/v1/runs", json={"job_id": job_id, "master_version_id": master["id"]}
    ).json()
    account.client.post(f"/api/v1/runs/{run['id']}/draft")
    account.run_url = f"/api/v1/runs/{run['id']}"  # type: ignore[attr-defined]
    artifact = render(account)["artifact"]
    assert artifact["layout"] == "reconstructed"
    assert any("reconstructed" in note for note in artifact["checks"]["notes"])
    assert artifact["passed"] is True and artifact["page_count"] == 1


def test_review_routes_are_private(demo_app, docx_bytes, base_settings):
    alice, run = drafted(demo_app, docx_bytes)
    artifact = render(alice)["artifact"]
    alice.client.post(
        f"{alice.run_url}/approve",
        json={
            "draft_version": artifact["draft_version"],
            "artifact_id": artifact["id"],
            "artifact_sha256": artifact["sha256"],
        },
    )
    approval_id = alice.client.get("/api/v1/approvals").json()[0]["id"]
    bob = new_account(demo_app)
    version = artifact["draft_version"]
    for method, path, body in [
        ("get", f"{alice.run_url}/draft/versions", None),
        ("get", f"{alice.run_url}/draft/versions/1", None),
        ("get", f"{alice.run_url}/chat", None),
        ("post", f"{alice.run_url}/render", None),
        (
            "post",
            f"{alice.run_url}/draft/edits",
            {"expected_version": version, "operations": [{"op": "update_summary", "text": None}]},
        ),
        ("post", f"{alice.run_url}/draft/restore", {"version": 1, "expected_version": version}),
        (
            "post",
            f"{alice.run_url}/approve",
            {
                "draft_version": version,
                "artifact_id": artifact["id"],
                "artifact_sha256": artifact["sha256"],
            },
        ),
        ("get", f"/api/v1/artifacts/{artifact['id']}/pdf", None),
        ("get", f"/api/v1/artifacts/{artifact['id']}/preview.png", None),
        ("get", f"/api/v1/approvals/{approval_id}/download", None),
    ]:
        call = getattr(bob.client, method)
        response = call(path, json=body) if body is not None else call(path)
        assert response.status_code == 404, (path, response.status_code, response.text[:200])
    assert bob.client.get("/api/v1/approvals").json() == []
    # Direct storage access with Bob's own token is refused by Storage policies.
    settings = base_settings
    pdf_path = f"{alice.id}/"
    listing = httpx.post(
        f"{settings.supabase_url}/storage/v1/object/list/artifacts",
        json={"prefix": pdf_path, "limit": 10},
        headers={
            "apikey": settings.supabase_anon_key.get_secret_value(),
            "Authorization": f"Bearer {bob.access_token}",
        },
        timeout=10,
    )
    assert listing.status_code in (200, 400, 403)
    if listing.status_code == 200:
        assert listing.json() == []  # Bob sees nothing in Alice's folder
